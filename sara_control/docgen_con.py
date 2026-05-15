"""
SARA Control — Document generation: DOCX, XLSX, PDF, image compose/edit.
Extracted from sara_controlgen1.py.
"""
import os
import importlib
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

try:
    from sara_control.loader_con import _load_core_module
    _core = _load_core_module()
    if _core is None:
        raise ImportError("Core module file not found")
    create_nbs_file = _core.create_nbs_file
except Exception:
    _core = None
    def create_nbs_file(project_name, relative_path, content, meta=None):
        return {"success": False, "error": "Core not loaded"}


def write_docx_con(
    con_output_path: str,
    con_paragraphs: List[str],
    con_title: Optional[str] = None,
    con_headings: Optional[List[Dict[str, Any]]] = None,
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Create a new DOCX file. Requires: python-docx"""
    try:
        docx = importlib.import_module("docx")
    except Exception:
        return {"success": False, "error": "python-docx not installed (pip install python-docx)"}
    try:
        doc = docx.Document()
        if con_title:
            doc.add_heading(con_title, level=0)
        for h in (con_headings or []):
            level = int(h.get("level", 1))
            text = str(h.get("text", ""))
            doc.add_heading(text, level=level)
        for para in (con_paragraphs or []):
            doc.add_paragraph(str(para))
        os.makedirs(os.path.dirname(os.path.abspath(con_output_path)), exist_ok=True)
        doc.save(con_output_path)
    except Exception as e:
        return {"success": False, "error": f"write_docx failed: {e}"}
    result: Dict[str, Any] = {"success": True, "path": con_output_path}
    if con_project:
        wrap = create_nbs_file(
            project_name=con_project,
            relative_path=os.path.join("artifacts", f"docx_written_{int(datetime.now(timezone.utc).timestamp())}_gen0_1.0_nbs.json"),
            content={"artifact_type": "docx_written", "path": con_output_path, "title": con_title, "paragraph_count": len(con_paragraphs or [])},
            meta={"nbs_type": "artifact_wrapper", "tags": ["docx", "written"], "source": "control", "project_name": con_project},
        )
        result["nbs_wrapper"] = wrap.get("reference", {}).get("file_path")
    return result


def edit_docx_con(
    con_docx_path: str,
    con_append_paragraphs: Optional[List[str]] = None,
    con_find_replace: Optional[List[Dict[str, str]]] = None,
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Append paragraphs or apply find/replace to an existing DOCX. Requires: python-docx"""
    try:
        docx = importlib.import_module("docx")
    except Exception:
        return {"success": False, "error": "python-docx not installed (pip install python-docx)"}
    if not os.path.exists(con_docx_path):
        return {"success": False, "error": f"File not found: {con_docx_path}"}
    try:
        doc = docx.Document(con_docx_path)
        replacements_done = 0
        for fr in (con_find_replace or []):
            find_str = str(fr.get("find", ""))
            repl_str = str(fr.get("replace", ""))
            if not find_str:
                continue
            for para in doc.paragraphs:
                for run in para.runs:
                    if find_str in run.text:
                        run.text = run.text.replace(find_str, repl_str)
                        replacements_done += 1
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                if find_str in run.text:
                                    run.text = run.text.replace(find_str, repl_str)
                                    replacements_done += 1
        appended = 0
        for para in (con_append_paragraphs or []):
            doc.add_paragraph(str(para))
            appended += 1
        doc.save(con_docx_path)
    except Exception as e:
        return {"success": False, "error": f"edit_docx failed: {e}"}
    result: Dict[str, Any] = {"success": True, "path": con_docx_path, "appended": appended, "replacements": replacements_done}
    if con_project:
        wrap = create_nbs_file(
            project_name=con_project,
            relative_path=os.path.join("artifacts", f"docx_edited_{int(datetime.now(timezone.utc).timestamp())}_gen0_1.0_nbs.json"),
            content={"artifact_type": "docx_edited", "path": con_docx_path, "appended": appended, "replacements": replacements_done},
            meta={"nbs_type": "artifact_wrapper", "tags": ["docx", "edited"], "source": "control", "project_name": con_project},
        )
        result["nbs_wrapper"] = wrap.get("reference", {}).get("file_path")
    return result


def write_xlsx_con(
    con_output_path: str,
    con_sheets: Dict[str, List[List[Any]]],
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Create a new XLSX workbook. Requires: openpyxl"""
    try:
        openpyxl = importlib.import_module("openpyxl")
    except Exception:
        return {"success": False, "error": "openpyxl not installed (pip install openpyxl)"}
    try:
        wb = openpyxl.Workbook()
        if wb.active and wb.active.title in ("Sheet", "Sheet1"):
            del wb[wb.active.title]
        for sheet_name, rows in (con_sheets or {}).items():
            ws = wb.create_sheet(title=str(sheet_name)[:31])
            for row in (rows or []):
                ws.append(list(row))
        os.makedirs(os.path.dirname(os.path.abspath(con_output_path)), exist_ok=True)
        wb.save(con_output_path)
        wb.close()
    except Exception as e:
        return {"success": False, "error": f"write_xlsx failed: {e}"}
    result: Dict[str, Any] = {"success": True, "path": con_output_path, "sheets": list((con_sheets or {}).keys())}
    if con_project:
        wrap = create_nbs_file(
            project_name=con_project,
            relative_path=os.path.join("artifacts", f"xlsx_written_{int(datetime.now(timezone.utc).timestamp())}_gen0_1.0_nbs.json"),
            content={"artifact_type": "xlsx_written", "path": con_output_path, "sheets": list((con_sheets or {}).keys())},
            meta={"nbs_type": "artifact_wrapper", "tags": ["xlsx", "written"], "source": "control", "project_name": con_project},
        )
        result["nbs_wrapper"] = wrap.get("reference", {}).get("file_path")
    return result


def edit_xlsx_con(
    con_xlsx_path: str,
    con_sheet_name: Optional[str] = None,
    con_append_rows: Optional[List[List[Any]]] = None,
    con_cell_updates: Optional[List[Dict[str, Any]]] = None,
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Append rows or update cells in an existing XLSX. Requires: openpyxl"""
    try:
        openpyxl = importlib.import_module("openpyxl")
    except Exception:
        return {"success": False, "error": "openpyxl not installed (pip install openpyxl)"}
    if not os.path.exists(con_xlsx_path):
        return {"success": False, "error": f"File not found: {con_xlsx_path}"}
    try:
        wb = openpyxl.load_workbook(con_xlsx_path)
        if con_sheet_name and con_sheet_name in wb.sheetnames:
            ws = wb[con_sheet_name]
        else:
            ws = wb.active
        appended_rows = 0
        for row in (con_append_rows or []):
            ws.append(list(row))
            appended_rows += 1
        cells_updated = 0
        for upd in (con_cell_updates or []):
            r = int(upd.get("row", 1))
            c = int(upd.get("col", 1))
            v = upd.get("value")
            ws.cell(row=r, column=c, value=v)
            cells_updated += 1
        wb.save(con_xlsx_path)
        wb.close()
    except Exception as e:
        return {"success": False, "error": f"edit_xlsx failed: {e}"}
    result: Dict[str, Any] = {"success": True, "path": con_xlsx_path, "appended_rows": appended_rows, "cells_updated": cells_updated}
    if con_project:
        wrap = create_nbs_file(
            project_name=con_project,
            relative_path=os.path.join("artifacts", f"xlsx_edited_{int(datetime.now(timezone.utc).timestamp())}_gen0_1.0_nbs.json"),
            content={"artifact_type": "xlsx_edited", "path": con_xlsx_path, "appended_rows": appended_rows, "cells_updated": cells_updated},
            meta={"nbs_type": "artifact_wrapper", "tags": ["xlsx", "edited"], "source": "control", "project_name": con_project},
        )
        result["nbs_wrapper"] = wrap.get("reference", {}).get("file_path")
    return result


def write_pdf_con(
    con_output_path: str,
    con_lines: List[str],
    con_title: Optional[str] = None,
    con_font_size: int = 12,
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Generate a new PDF file. Tries fpdf2 first, falls back to reportlab."""
    os.makedirs(os.path.dirname(os.path.abspath(con_output_path)), exist_ok=True)
    backend = None
    try:
        fpdf_mod = importlib.import_module("fpdf")
        pdf = fpdf_mod.FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        if con_title:
            pdf.set_font("Helvetica", style="B", size=int(con_font_size) + 6)
            pdf.multi_cell(0, 10, txt=str(con_title))
            pdf.ln(4)
        pdf.set_font("Helvetica", size=int(con_font_size))
        for line in (con_lines or []):
            pdf.multi_cell(0, 8, txt=str(line))
        pdf.output(con_output_path)
        backend = "fpdf2"
    except Exception as fpdf_err:
        try:
            rl_canvas = importlib.import_module("reportlab.pdfgen.canvas")
            rl_lib = importlib.import_module("reportlab.lib.pagesizes")
            c = rl_canvas.canvas.Canvas(con_output_path, pagesize=rl_lib.letter)
            width, height = rl_lib.letter
            x = 50
            y = height - 60
            fsize = int(con_font_size)
            if con_title:
                c.setFont("Helvetica-Bold", fsize + 6)
                c.drawString(x, y, str(con_title))
                y -= fsize + 10
            c.setFont("Helvetica", fsize)
            for line in (con_lines or []):
                if y < 60:
                    c.showPage()
                    y = height - 60
                    c.setFont("Helvetica", fsize)
                c.drawString(x, y, str(line))
                y -= fsize + 4
            c.save()
            backend = "reportlab"
        except Exception as rl_err:
            return {"success": False, "error": f"No PDF backend available. fpdf2: {fpdf_err}  reportlab: {rl_err}"}
    result: Dict[str, Any] = {"success": True, "path": con_output_path, "backend": backend}
    if con_project:
        wrap = create_nbs_file(
            project_name=con_project,
            relative_path=os.path.join("artifacts", f"pdf_written_{int(datetime.now(timezone.utc).timestamp())}_gen0_1.0_nbs.json"),
            content={"artifact_type": "pdf_written", "path": con_output_path, "backend": backend, "line_count": len(con_lines or [])},
            meta={"nbs_type": "artifact_wrapper", "tags": ["pdf", "written"], "source": "control", "project_name": con_project},
        )
        result["nbs_wrapper"] = wrap.get("reference", {}).get("file_path")
    return result


def compose_image_con(
    con_output_path: str,
    con_width: int,
    con_height: int,
    con_bg_color: Optional[str] = "white",
    con_draw_ops: Optional[List[Dict[str, Any]]] = None,
    con_base_image_path: Optional[str] = None,
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Create or composite a PNG/JPEG image using PIL (Pillow)."""
    try:
        pil_image = importlib.import_module("PIL.Image")
        pil_draw = importlib.import_module("PIL.ImageDraw")
        pil_font = importlib.import_module("PIL.ImageFont")
    except Exception:
        return {"success": False, "error": "Pillow not installed (pip install Pillow)"}
    try:
        if con_base_image_path and os.path.exists(con_base_image_path):
            img = pil_image.open(con_base_image_path).convert("RGBA")
            img = img.resize((int(con_width), int(con_height)))
        else:
            bg = con_bg_color or "white"
            img = pil_image.new("RGBA", (int(con_width), int(con_height)), bg)
        draw = pil_draw.ImageDraw(img)
        for op_dict in (con_draw_ops or []):
            op = str(op_dict.get("op", "")).lower()
            if op == "rectangle":
                draw.rectangle(op_dict.get("xy"), fill=op_dict.get("fill"), outline=op_dict.get("outline"))
            elif op == "ellipse":
                draw.ellipse(op_dict.get("xy"), fill=op_dict.get("fill"), outline=op_dict.get("outline"))
            elif op == "line":
                draw.line(op_dict.get("xy"), fill=op_dict.get("fill"), width=int(op_dict.get("width", 1)))
            elif op == "text":
                fsize = int(op_dict.get("font_size", 14))
                try:
                    font = pil_font.truetype("arial.ttf", fsize)
                except Exception:
                    font = pil_font.load_default()
                draw.text(op_dict.get("xy", (0, 0)), str(op_dict.get("text", "")), fill=op_dict.get("fill", "black"), font=font)
            elif op == "paste":
                src = str(op_dict.get("src", ""))
                if os.path.exists(src):
                    overlay = pil_image.open(src).convert("RGBA")
                    img.paste(overlay, op_dict.get("xy", (0, 0)), overlay)
        os.makedirs(os.path.dirname(os.path.abspath(con_output_path)), exist_ok=True)
        ext = os.path.splitext(con_output_path)[1].lower()
        if ext in (".jpg", ".jpeg"):
            img.convert("RGB").save(con_output_path)
        else:
            img.save(con_output_path)
    except Exception as e:
        return {"success": False, "error": f"compose_image failed: {e}"}
    result: Dict[str, Any] = {"success": True, "path": con_output_path}
    if con_project:
        wrap = create_nbs_file(
            project_name=con_project,
            relative_path=os.path.join("artifacts", f"image_composed_{int(datetime.now(timezone.utc).timestamp())}_gen0_1.0_nbs.json"),
            content={"artifact_type": "image_composed", "path": con_output_path, "width": con_width, "height": con_height, "ops": len(con_draw_ops or [])},
            meta={"nbs_type": "artifact_wrapper", "tags": ["image", "composed"], "source": "control", "project_name": con_project},
        )
        result["nbs_wrapper"] = wrap.get("reference", {}).get("file_path")
    return result


def write_xlsx_formula_con(
    con_output_path: str,
    con_sheets: Dict[str, List[List[Any]]],
    con_column_widths: Optional[Dict[str, List[int]]] = None,
    con_bold_rows: Optional[List[int]] = None,
    con_freeze_pane: Optional[str] = None,
    con_auto_filter: bool = False,
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Create an XLSX with formula support, formatting, freeze panes, and auto-filter."""
    try:
        openpyxl = importlib.import_module("openpyxl")
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter
    except Exception:
        return {"success": False, "error": "openpyxl not installed"}
    try:
        wb = openpyxl.Workbook()
        if wb.active and wb.active.title in ("Sheet", "Sheet1"):
            del wb[wb.active.title]
        for sheet_name, rows in (con_sheets or {}).items():
            ws = wb.create_sheet(title=str(sheet_name)[:31])
            for r_idx, row in enumerate(rows or [], start=1):
                for c_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    if isinstance(val, str) and val.startswith("="):
                        cell.value = val
                    else:
                        cell.value = val
                    if r_idx in (con_bold_rows or []):
                        cell.font = Font(bold=True)
            widths = (con_column_widths or {}).get(sheet_name, [])
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            if con_freeze_pane:
                ws.freeze_panes = con_freeze_pane
            if con_auto_filter and ws.max_row and ws.max_column:
                ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        os.makedirs(os.path.dirname(os.path.abspath(con_output_path)), exist_ok=True)
        wb.save(con_output_path)
        wb.close()
    except Exception as e:
        return {"success": False, "error": f"write_xlsx_formula failed: {e}"}
    return {"success": True, "path": con_output_path}


def write_xlsx_chart_con(
    con_xlsx_path: str,
    con_chart_type: str = "bar",
    con_data_range: str = "A1:B10",
    con_title: str = "Chart",
    con_sheet_name: Optional[str] = None,
    con_chart_sheet: str = "Chart",
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Add a chart to an existing XLSX workbook."""
    try:
        openpyxl = importlib.import_module("openpyxl")
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.utils import range_boundaries
    except Exception:
        return {"success": False, "error": "openpyxl not installed"}
    if not os.path.exists(con_xlsx_path):
        return {"success": False, "error": f"File not found: {con_xlsx_path}"}
    try:
        wb = openpyxl.load_workbook(con_xlsx_path)
        ws = wb[con_sheet_name] if con_sheet_name and con_sheet_name in wb.sheetnames else wb.active

        min_col, min_row, max_col, max_row = range_boundaries(con_data_range)

        chart_classes = {"bar": BarChart, "line": LineChart, "pie": PieChart}
        ChartClass = chart_classes.get(con_chart_type.lower(), BarChart)
        chart = ChartClass()
        chart.title = con_title
        chart.width = 18
        chart.height = 12

        cats = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
        for col_idx in range(min_col + 1, max_col + 1):
            data = Reference(ws, min_col=col_idx, min_row=min_row, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        if con_chart_sheet in wb.sheetnames:
            chart_ws = wb[con_chart_sheet]
        else:
            chart_ws = wb.create_sheet(title=con_chart_sheet)
        chart_ws.add_chart(chart, "A1")
        wb.save(con_xlsx_path)
        wb.close()
    except Exception as e:
        return {"success": False, "error": f"chart creation failed: {e}"}
    return {"success": True, "path": con_xlsx_path, "chart_type": con_chart_type}


def edit_image_con(
    con_image_path: str,
    con_resize: Optional[tuple] = None,
    con_crop: Optional[tuple] = None,
    con_overlay_path: Optional[str] = None,
    con_overlay_pos: Optional[tuple] = None,
    con_annotate_text: Optional[List[Dict[str, Any]]] = None,
    con_project: Optional[str] = None,
) -> Dict[str, Any]:
    """Edit an existing image: resize, crop, overlay, annotate. Requires: Pillow"""
    try:
        pil_image = importlib.import_module("PIL.Image")
        pil_draw = importlib.import_module("PIL.ImageDraw")
        pil_font = importlib.import_module("PIL.ImageFont")
    except Exception:
        return {"success": False, "error": "Pillow not installed (pip install Pillow)"}
    if not os.path.exists(con_image_path):
        return {"success": False, "error": f"Image not found: {con_image_path}"}
    try:
        img = pil_image.open(con_image_path).convert("RGBA")
        if con_crop:
            img = img.crop(con_crop)
        if con_resize:
            img = img.resize(con_resize, pil_image.Resampling.LANCZOS)
        if con_overlay_path and os.path.exists(con_overlay_path):
            overlay = pil_image.open(con_overlay_path).convert("RGBA")
            pos = tuple(con_overlay_pos) if con_overlay_pos else (0, 0)
            img.paste(overlay, pos, overlay)
        draw = pil_draw.ImageDraw(img)
        for ann in (con_annotate_text or []):
            fsize = int(ann.get("font_size", 14))
            try:
                font = pil_font.truetype("arial.ttf", fsize)
            except Exception:
                font = pil_font.load_default()
            draw.text(ann.get("xy", (0, 0)), str(ann.get("text", "")), fill=ann.get("fill", "black"), font=font)
        ext = os.path.splitext(con_image_path)[1].lower()
        if ext in (".jpg", ".jpeg"):
            img.convert("RGB").save(con_image_path)
        else:
            img.save(con_image_path)
    except Exception as e:
        return {"success": False, "error": f"edit_image failed: {e}"}
    result: Dict[str, Any] = {"success": True, "path": con_image_path}
    if con_project:
        wrap = create_nbs_file(
            project_name=con_project,
            relative_path=os.path.join("artifacts", f"image_edited_{int(datetime.now(timezone.utc).timestamp())}_gen0_1.0_nbs.json"),
            content={"artifact_type": "image_edited", "path": con_image_path, "resized": con_resize is not None, "cropped": con_crop is not None, "annotated": len(con_annotate_text or [])},
            meta={"nbs_type": "artifact_wrapper", "tags": ["image", "edited"], "source": "control", "project_name": con_project},
        )
        result["nbs_wrapper"] = wrap.get("reference", {}).get("file_path")
    return result
