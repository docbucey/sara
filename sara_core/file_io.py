"""CORE File I/O: readers and writers for all supported file formats."""
import os
import json
import struct
from typing import Any, Dict, List, Optional

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import docx  # python-docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pptx  # python-pptx
except ImportError:
    pptx = None


# Image/Graphics File IO Handlers (JPG, PNG, GIF, etc.)
# =========================================================
def read_image_file(path):
    if not Image:
        return {"success": False, "error": "Pillow not installed"}
    try:
        img = Image.open(path)
        info = {
            "format": img.format,
            "mode": img.mode,
            "size": img.size
        }
        return {"success": True, "info": info}
    except Exception as e:
        return {"success": False, "error": str(e)}

def write_image_file(path, data, format=None):
    if not Image:
        return {"success": False, "error": "Pillow not installed"}
    try:
        img = Image.fromarray(data) if hasattr(data, 'shape') else Image.open(data)
        img.save(path, format=format)
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}

def resize_image_file(path, output_path, size):
    if not Image:
        return {"success": False, "error": "Pillow not installed"}
    try:
        img = Image.open(path)
        img = img.resize(size)
        img.save(output_path)
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}


def read_video_file(path):
    """Read video asset metadata with optional frame-level detail when cv2 is available."""
    import os

    if not os.path.exists(path):
        return {"success": False, "error": "file not found"}

    ext = os.path.splitext(path)[1].lower()
    if ext not in {".mp4", ".mov", ".avi", ".mkv", ".webm"}:
        return {"success": False, "error": f"unsupported video format: {ext}"}

    info = {
        "path": path,
        "ext": ext,
        "size_bytes": os.path.getsize(path),
    }

    try:
        import cv2  # optional

        cap = cv2.VideoCapture(path)
        if cap.isOpened():
            fps = float(cap.get(cv2.CAP_PROP_FPS) or 0.0)
            frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
            width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH) or 0)
            height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT) or 0)
            duration = float(frame_count / fps) if fps > 0 else None
            info.update(
                {
                    "fps": fps,
                    "frame_count": frame_count,
                    "width": width,
                    "height": height,
                    "duration_seconds": duration,
                    "metadata_mode": "cv2",
                }
            )
        cap.release()
    except Exception:
        info["metadata_mode"] = "basic"

    return {"success": True, "info": info}


def read_audio_file(path):
    """Read audio asset metadata for common office-suite media lanes."""
    import os

    if not os.path.exists(path):
        return {"success": False, "error": "file not found"}

    ext = os.path.splitext(path)[1].lower()
    if ext not in {".mp3", ".wav", ".flac", ".ogg", ".m4a", ".aac"}:
        return {"success": False, "error": f"unsupported audio format: {ext}"}

    info = {
        "path": path,
        "ext": ext,
        "size_bytes": os.path.getsize(path),
        "metadata_mode": "basic",
    }

    if ext == ".wav":
        try:
            import wave

            with wave.open(path, "rb") as wav:
                rate = int(wav.getframerate())
                frames = int(wav.getnframes())
                channels = int(wav.getnchannels())
                width = int(wav.getsampwidth())
                info.update(
                    {
                        "sample_rate": rate,
                        "channels": channels,
                        "sample_width_bytes": width,
                        "duration_seconds": float(frames / rate) if rate > 0 else None,
                        "metadata_mode": "wave",
                    }
                )
        except Exception:
            pass

    return {"success": True, "info": info}


def read_3d_file(path):
    """Read 3D asset metadata for OBJ/STL files used by modern graphics workflows."""
    import os
    import struct

    if not os.path.exists(path):
        return {"success": False, "error": "file not found"}

    ext = os.path.splitext(path)[1].lower()
    if ext not in {".obj", ".stl"}:
        return {"success": False, "error": f"unsupported 3d format: {ext}"}

    info = {
        "path": path,
        "ext": ext,
        "size_bytes": os.path.getsize(path),
    }

    try:
        if ext == ".obj":
            vertices = 0
            texcoords = 0
            normals = 0
            faces = 0
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                for line in f:
                    s = line.lstrip()
                    if s.startswith("v "):
                        vertices += 1
                    elif s.startswith("vt "):
                        texcoords += 1
                    elif s.startswith("vn "):
                        normals += 1
                    elif s.startswith("f "):
                        faces += 1
            info.update(
                {
                    "vertices": vertices,
                    "texcoords": texcoords,
                    "normals": normals,
                    "faces": faces,
                    "metadata_mode": "obj_text_scan",
                }
            )
        else:
            with open(path, "rb") as f:
                header = f.read(84)
            if len(header) >= 84:
                tri_count = struct.unpack("<I", header[80:84])[0]
                expected = 84 + tri_count * 50
                if expected == info["size_bytes"]:
                    info.update({"triangles": int(tri_count), "format_hint": "binary", "metadata_mode": "stl_binary_header"})
                else:
                    facets = 0
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        for line in f:
                            if line.lstrip().lower().startswith("facet normal"):
                                facets += 1
                    info.update({"triangles": facets, "format_hint": "ascii", "metadata_mode": "stl_ascii_scan"})
    except Exception as e:
        return {"success": False, "error": str(e)}

    return {"success": True, "info": info}


def read_graphics_asset_file(path):
    """Unified reader for image/video/audio/3D assets in the office graphics lane."""
    import os

    ext = os.path.splitext(path)[1].lower()
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"}:
        return read_image_file(path)
    if ext in {".mp4", ".mov", ".avi", ".mkv", ".webm"}:
        return read_video_file(path)
    if ext in {".mp3", ".wav", ".flac", ".ogg", ".m4a", ".aac"}:
        return read_audio_file(path)
    if ext in {".stl", ".obj"}:
        return read_3d_file(path)
    return {"success": False, "error": f"unsupported graphics asset type: {ext}"}


# =========================================================
# Office File IO Handlers (Word, Excel, PowerPoint)
# All actions must be routed through control for enforcement/execution.
# =========================================================
def read_word_docx(path):
    if not docx:
        return {"success": False, "error": "python-docx not installed"}
    try:
        doc = docx.Document(path)
        text = "\n".join([p.text for p in doc.paragraphs])
        return {"success": True, "text": text}
    except Exception as e:
        return {"success": False, "error": str(e)}

def write_word_docx(path, paragraphs):
    if not docx:
        return {"success": False, "error": "python-docx not installed"}
    try:
        doc = docx.Document()
        for para in paragraphs:
            doc.add_paragraph(para)
        doc.save(path)
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}

def read_excel_xlsx(path):
    if not openpyxl:
        return {"success": False, "error": "openpyxl not installed"}
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        data = {}
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            data[sheet] = [[cell.value for cell in row] for row in ws.iter_rows()]
        return {"success": True, "sheets": data}
    except Exception as e:
        return {"success": False, "error": str(e)}

def write_excel_xlsx(path, sheets_data):
    if not openpyxl:
        return {"success": False, "error": "openpyxl not installed"}
    try:
        wb = openpyxl.Workbook()
        for idx, (sheet, rows) in enumerate(sheets_data.items()):
            ws = wb.create_sheet(title=sheet) if idx > 0 else wb.active
            ws.title = sheet
            for row in rows:
                ws.append(row)
        wb.save(path)
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}

def read_powerpoint_pptx(path):
    if not pptx:
        return {"success": False, "error": "python-pptx not installed"}
    try:
        prs = pptx.Presentation(path)
        slides = []
        for slide in prs.slides:
            text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
            slides.append("\n".join(text))
        return {"success": True, "slides": slides}
    except Exception as e:
        return {"success": False, "error": str(e)}

def write_powerpoint_pptx(path, slides_data):
    if not pptx:
        return {"success": False, "error": "python-pptx not installed"}
    try:
        prs = pptx.Presentation()
        for slide_text in slides_data:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            shapes = slide.shapes
            shapes.title.text = slide_text[0] if slide_text else ""
            for para in slide_text[1:]:
                shapes.placeholders[1].text += "\n" + para
        prs.save(path)
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}
