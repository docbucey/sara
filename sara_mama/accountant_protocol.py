"""
AccountantProtocol: Excel/Google Sheets style spreadsheet protocol for Mama
Implements spreadsheet validation and automation logic using pandas and openpyxl.
"""

from typing import Optional, List, Dict, Any

try:
    import pandas as pd
    import openpyxl
    _EXCEL_OK = True
except Exception:
    _EXCEL_OK = False

class AccountantProtocol:
    """
    Protocol for validating and automating spreadsheet documents (Excel, Google Sheets style).
    """
    def __init__(self):
        if not _EXCEL_OK:
            raise ImportError("pandas and openpyxl are required for AccountantProtocol")

    def validate_spreadsheet(self, excel_path: str, required_sheets: Optional[List[str]] = None, required_columns: Optional[List[str]] = None) -> Dict[str, Any]:
        """
        Validate that the spreadsheet contains all required sheets and columns.
        """
        if not _EXCEL_OK:
            return {"status": "ERROR", "error": "pandas/openpyxl not available"}
        try:
            wb = openpyxl.load_workbook(excel_path, read_only=True)
            sheet_names = wb.sheetnames
            missing_sheets = [s for s in (required_sheets or []) if s not in sheet_names]
            missing_columns = {}
            for sheet in (required_sheets or sheet_names):
                if sheet in sheet_names:
                    df = pd.read_excel(excel_path, sheet_name=sheet)
                    cols = list(df.columns)
                    missing = [c for c in (required_columns or []) if c not in cols]
                    if missing:
                        missing_columns[sheet] = missing
            return {
                "status": "PASS" if not missing_sheets and not missing_columns else "FAIL",
                "sheet_names": sheet_names,
                "missing_sheets": missing_sheets,
                "missing_columns": missing_columns,
                "required_sheets": required_sheets or [],
                "required_columns": required_columns or [],
                "excel_path": excel_path,
            }
        except Exception as e:
            return {"status": "ERROR", "error": str(e), "excel_path": excel_path}

    def automate_formatting(self, excel_path: str, output_path: Optional[str] = None) -> Dict[str, Any]:
        """
        Example automation: set all header rows to bold in all sheets.
        """
        if not _EXCEL_OK:
            return {"status": "ERROR", "error": "pandas/openpyxl not available"}
        try:
            wb = openpyxl.load_workbook(excel_path)
            changed = 0
            for ws in wb.worksheets:
                for cell in ws[1]:  # First row (header)
                    cell.font = openpyxl.styles.Font(bold=True)
                    changed += 1
            out_path = output_path or excel_path
            wb.save(out_path)
            return {"status": "PASS", "changed_headers": changed, "output_path": out_path}
        except Exception as e:
            return {"status": "ERROR", "error": str(e), "excel_path": excel_path}
