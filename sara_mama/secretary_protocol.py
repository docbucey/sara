"""
SecretaryProtocol: MS Word / Google Docs style word processor protocol for Mama
Implements document validation and automation logic using python-docx.
"""

from typing import Optional, List, Dict, Any

try:
    import docx  # python-docx
    _DOCX_OK = True
except Exception:
    _DOCX_OK = False

class SecretaryProtocol:
    """
    Protocol for validating and automating word processor documents (MS Word, Google Docs style).
    """
    def __init__(self):
        if not _DOCX_OK:
            raise ImportError("python-docx is required for SecretaryProtocol")

    def validate_document(self, docx_path: str, required_sections: Optional[List[str]] = None) -> Dict[str, Any]:
        """
        Validate that the document contains all required sections (by heading text).
        """
        if not _DOCX_OK:
            return {"status": "ERROR", "error": "python-docx not available"}
        try:
            doc = docx.Document(docx_path)
            headings = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading")]
            missing = []
            for section in (required_sections or []):
                if section not in headings:
                    missing.append(section)
            return {
                "status": "PASS" if not missing else "FAIL",
                "headings": headings,
                "missing_sections": missing,
                "required_sections": required_sections or [],
                "docx_path": docx_path,
            }
        except Exception as e:
            return {"status": "ERROR", "error": str(e), "docx_path": docx_path}

    def automate_formatting(self, docx_path: str, output_path: Optional[str] = None) -> Dict[str, Any]:
        """
        Example automation: ensure all headings are bold and font size 14pt.
        """
        if not _DOCX_OK:
            return {"status": "ERROR", "error": "python-docx not available"}
        try:
            doc = docx.Document(docx_path)
            changed = 0
            for p in doc.paragraphs:
                if p.style.name.startswith("Heading"):
                    for run in p.runs:
                        run.bold = True
                        run.font.size = docx.shared.Pt(14)
                    changed += 1
            out_path = output_path or docx_path
            doc.save(out_path)
            return {"status": "PASS", "changed_headings": changed, "output_path": out_path}
        except Exception as e:
            return {"status": "ERROR", "error": str(e), "docx_path": docx_path}
